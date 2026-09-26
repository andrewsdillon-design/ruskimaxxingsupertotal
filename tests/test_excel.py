import csv
import re
import shutil
import subprocess
import zipfile
from datetime import date

import pytest
from openpyxl import load_workbook

from ruskimaxxing.excel import build_workbook
from ruskimaxxing.exercises import CATALOG

DATA = {"units": "lb", "increment": 5, "start": date(2026, 10, 5), "height": 70, "bodyweight_start": 178,
        "bodyfat_start": 19, "bodyfat_method": "Bod Pod",
        "baseline": [("Squat", 250, 5), ("Bench Press", 185, 1), ("2-Board Press", 200, 5), ("Box Jump", 24, 1),
                     ("Broad Jump", 80, 1)],
        "log": [(date(2026, 12, 26), "Squat", 300, 1, "test"), (date(2026, 11, 1), "1-Board Press", 190, 3, "")],
        "bodyweight": {1: 180, 13: 185},
        "bodyfat": {1: (date(2026, 10, 6), 18.5, "InBody", 180)}}

MONTHS = [f"Month {m:02d}" for m in range(1, 14)]


def rows_named(ws, name):
    return [r for r in range(1, ws.max_row + 1) if ws.cell(r, 3).value == name]


def test_sheet_order_and_layout(tmp_path):
    wb = load_workbook(build_workbook(tmp_path / "p.xlsx", DATA))
    assert wb.sheetnames == ["Start Here", "Baseline", *MONTHS, "Log", "PRs", "Maxes", "Body", "Progress",
                             "Sets", "Lists"]
    assert wb["Sets"].sheet_state == wb["Lists"].sheet_state == "hidden"
    start = wb["Start Here"]
    assert "malchiki" in start["A3"].value and "5 lb" in start["A3"].value
    labels = [start.cell(r, 1).value for r in range(5, 15)]
    assert any("HEIGHT" in l for l in labels) and any("BODYWEIGHT" in l for l in labels)
    m1 = wb["Month 01"]
    assert m1["D3"].value.date() == date(2026, 10, 6) and m1["F3"].value == 18.5          # monthly body fat
    week1 = next(r for r in range(1, 30) if str(m1.cell(r, 1).value).startswith("WEEK 1 "))
    assert m1.cell(week1, 9).value == 180                                          # weekly bodyweight
    squat = rows_named(m1, "Squat")[0]
    assert [m1.cell(squat + k, 3).value.strip() for k in (1, 2, 3)] == ["Set 1", "Set 2", "Set 3"]
    assert m1.cell(squat + 1, 10).value == 6                                      # reps pre-filled
    assert m1.cell(squat + 1, 9).value == f'=IF(H{squat}="","",H{squat})'         # weight pre-filled


def test_no_external_links_and_schema_font_order(tmp_path):
    path = build_workbook(tmp_path / "p.xlsx", DATA)
    with zipfile.ZipFile(path) as z:
        rels = [n for n in z.namelist() if n.endswith(".rels")]
        assert not any('TargetMode="External"' in z.read(n).decode() for n in rels)
        styles = z.read("xl/styles.xml").decode()
    for font in re.findall(r"<font>(.*?)</font>", styles):  # color must never come before sz / u
        tags = re.findall(r"<(\w+)", font)
        if "color" in tags and "sz" in tags:
            assert tags.index("sz") < tags.index("color")
    wb = load_workbook(path)
    links = [c.hyperlink for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.hyperlink]
    assert len(links) > 50 and all(h.location and not h.target for h in links)


def test_every_catalog_exercise_has_a_pr_row(tmp_path):
    wb = load_workbook(build_workbook(tmp_path / "p.xlsx"))
    assert [c.value for c in wb["PRs"]["A"][1:len(CATALOG) + 1]] == list(CATALOG)


def _recalc(tmp_path, edit=None):
    path = build_workbook(tmp_path / "p.xlsx", DATA)
    if edit:
        wb = load_workbook(path)
        edit(wb)
        wb.save(path)
    subprocess.run(["soffice", "--headless", "--convert-to",
                    "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1",
                    "--outdir", str(tmp_path), str(path)], check=True, capture_output=True, timeout=300)

    def sheet(name):
        with open(tmp_path / f"p-{name}.csv", newline="", encoding="utf-8") as f:
            return list(csv.reader(f))
    return sheet


needs_lo = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")


@needs_lo
def test_formulas_calculate_without_errors_or_cycles(tmp_path):
    sheet = _recalc(tmp_path)
    for name in ["Start Here", "Baseline", *MONTHS, "PRs", "Maxes", "Body", "Progress"]:
        text = "\n".join(",".join(r) for r in sheet(name))
        assert not any(err in text for err in ("#NAME", "#VALUE", "#REF", "Err:")), name  # Err:522 = circular

    maxes = {r[0]: r for r in sheet("Maxes")}
    assert maxes["Squat"][4:6] == ["291.7", "300"]         # 250x5 baseline, then a 300 single before cycle 2
    assert maxes["Box Squat"][4] == "262.5"                # estimated: squat x 0.9
    assert maxes["2-Board Press"][4] == "233.3"            # its own logged 5RM
    assert maxes["Box Jump"][4] == "24"
    month4 = sheet("Month 04")
    squat = next(i for i, r in enumerate(month4) if r[2] == "Squat")
    assert month4[squat][6:8] == ["300", "210"]            # cycle 2 uses the new max: 70% of 300
    assert month4[squat + 1][8:10] == ["210", "6"]         # set row pre-filled
    start = "\n".join(",".join(r) for r in sheet("Start Here"))
    assert "Intermediate" in start                         # 24 in box / 80 in broad jump at 70 in tall
    assert sheet("Body")[3][10] == "144.2"                 # lean mass from the intake (178 lb, 19%)


@needs_lo
def test_only_done_sets_count(tmp_path):
    def log(wb):
        ws = wb["Month 01"]
        squat = rows_named(ws, "Squat")[0]
        ws.cell(squat + 1, 9).value = 320                  # done -> counts
        ws.cell(squat + 1, 12).value = "✓"
        ws.cell(squat + 2, 9).value = 400                  # NOT done -> must not count
        box = rows_named(ws, "Box Jump")[0]
        ws.cell(box + 1, 9).value = 30
        ws.cell(box + 1, 12).value = "✓"

    sheet = _recalc(tmp_path, log)
    prs = {r[0]: r for r in sheet("PRs")}
    assert prs["Squat"][3] == "320" and prs["Box Jump"][3] == "30"
    month1 = sheet("Month 01")
    squat = next(i for i, r in enumerate(month1) if r[2] == "Squat")
    assert month1[squat + 1][13] == "PR!" and month1[squat + 2][12] == ""
    assert sheet("Maxes")[1][5] == "384"  # cycle 2 squat TM = 320 x 6 -> e1RM 384 (the undone 400 ignored)
