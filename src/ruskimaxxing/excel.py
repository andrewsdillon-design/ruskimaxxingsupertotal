"""Build the standalone Excel workbook.

Everything is driven by formulas, so the file works on its own in Excel 2019+,
Microsoft 365, LibreOffice or Google Sheets. Sheet order:

  Start Here      intake, starting maxes, PR board, box jump standards,
                  Prilepin's chart, links to every week, instructions
  Week 00-52      one sheet per week: weights, log actual sets, weekly bodyweight
  Log             any extra sets (off-program work, extra variations, re-tests)
  PRs             best e1RM and 1-12 rep maxes for every movement and variation
  Maxes           training max per exercise per cycle (best e1RM before the cycle)
  Body            weekly bodyweight (from the week sheets), monthly body fat, test guide
  Progress        charts: strength vs bodyweight, bodyweight, body fat / lean mass
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ruskimaxxing.edition import app_name
from ruskimaxxing.exercises import BOX_JUMP_STANDARDS, CATALOG, MAIN
from ruskimaxxing.prilepin import ZONES
from ruskimaxxing.program import (CYCLE_WEEKS, CYCLES, WEEKS, bodyfat_week, build_program,
                                  next_monday, week_label)
from ruskimaxxing.tracking import BODYFAT_GUIDE, BODYFAT_METHODS, REP_MAX_COUNTS

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=13, color="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
DAY_FILL = PatternFill("solid", fgColor="EEF3F8")
LINK_FONT = Font(color="0563C1", underline="single")
PR_FONT = Font(bold=True, color="B00020")
PHASE_FILLS = {
    "Baseline": "DDE3F7", "Accumulation": "DCEFDC", "Transmutation": "FCE9D2", "Realization": "F8D7D7",
    "Deload": "E6E6E6", "Taper": "EDE0F5", "Test": "DDE3F7", "Transition": "E0F2F1",
}
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HIDE_ZERO = '#,##0.#;-#,##0.#;""'

INTAKE = [  # (label, key, default) - rows 4..13 of Start Here, column B
    ("Name (optional)", "name", None),
    ("Units", "units", "lb"),
    ("Round weights to nearest", "increment", 5),
    ("Start date (a Monday)", "start", None),
    ("Height (inches if lb, cm if kg)", "height", None),
    ("Starting bodyweight", "bodyweight_start", None),
    ("Starting body fat % (optional)", "bodyfat_start", None),
    ("Body fat method", "bodyfat_method", None),
    ("Age (optional)", "age", None),
    ("Sex (optional)", "sex", None),
]
SH = "'Start Here'"


def _intake_ref(key: str) -> str:
    return f"{SH}!$B${4 + [k for _, k, _ in INTAKE].index(key)}"


UNITS, ROUNDING, START, HEIGHT = (_intake_ref(k) for k in ("units", "increment", "start", "height"))
START_BW, START_BF, BF_METHOD = (_intake_ref(k) for k in ("bodyweight_start", "bodyfat_start", "bodyfat_method"))
BASE_FIRST, BASE_ROWS = 17, 40
LOG_ROWS = 1000
WEEK_FIRST_ROW = 6

INSTRUCTIONS = [
    "How to use this workbook",
    "1. Fill in the intake (yellow cells, top left): units, rounding, the Monday you start, height, bodyweight.",
    "2. Starting maxes - EITHER enter what you know in the Starting maxes table (any movement or variation,",
    "   weight and reps: 1 for a true max, or e.g. 185 x 5), OR run Week 00 - a test week - and log it there.",
    "   Variations you haven't done yet are estimated from the main lift until you log them.",
    "3. Each week has its own sheet (Week 01 ... Week 52; links above). Weights fill in automatically.",
    "   Log the weight and reps of your top set in the yellow columns - PRs are flagged in red.",
    "   For jumps, log the box height / jump distance in the 'Actual' column.",
    "4. Enter your bodyweight once a week at the top of each week sheet (same day, morning, before eating).",
    "5. Get body fat measured once a month (Bod Pod, InBody or hydrostatic) - the week sheet reminds you;",
    "   enter it on the Body sheet, which explains how to prepare for each kind of test.",
    "6. Every 12th week (12, 24, 36, 48) is a test week. Those numbers become next cycle's training maxes.",
    "7. The PR board shows your main lifts; the PRs sheet has every movement and variation.",
    "   The Progress sheet charts strength next to bodyweight and body fat.",
    "8. Train 3 days a week (Mon/Wed/Fri). Warm up before every main lift: empty bar x 10, then 3-4 sets up.",
    "",
    "Eat enough to grow: a small calorie surplus and ~0.7-1 g protein per lb bodyweight (1.6-2.2 g/kg). Sleep 7-9 h.",
    "Free for everyone - MIT license. Not medical advice; check with a doctor before starting.",
]


def week_sheet(week: int) -> str:
    return f"Week {week:02d}"


def _header(ws, row, headers, first_col=1):
    for col, text in enumerate(headers, start=first_col):
        c = ws.cell(row, col, text)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths, first_col=1):
    for i, w in enumerate(widths, start=first_col):
        ws.column_dimensions[get_column_letter(i)].width = w


def _input(cell, value=None):
    if value is not None:
        cell.value = value
    cell.fill, cell.border = INPUT_FILL, BOX
    return cell


def _link(cell, text, sheet):
    cell.value = text
    cell.hyperlink = f"#'{sheet}'!A1"
    cell.font = LINK_FONT
    return cell


def _e1rm(w, r):
    return f'=IF(OR({w}="",{r}=""),"",IF({r}=1,{w},ROUND({w}*(1+{r}/30),1)))'


class _Book:
    def __init__(self, data: dict | None):
        self.data = data or {}
        self.wb = Workbook()
        self.lists = self.wb.active
        self.lists.title = "Lists"
        names = list(CATALOG)
        for i, name in enumerate(names, start=1):
            self.lists.cell(i, 1, name)
        for i, m in enumerate(BODYFAT_METHODS, start=1):
            self.lists.cell(i, 2, m)
        self.lists.sheet_state = "hidden"
        self.ex_list = f"Lists!$A$1:$A${len(names)}"
        self.method_list = f"Lists!$B$1:$B${len(BODYFAT_METHODS)}"
        self.row_of = {name: 2 + i for i, name in enumerate(names)}  # same row on PRs and Maxes
        self.sets_refs = []
        self.log_rows = max(LOG_ROWS, len(self.data.get("log", [])) + 200)
        self.sessions = build_program()
        self.bf_weeks = {bodyfat_week(m): m for m in range(1, 13)}
        self.jump_targets = {}

    @property
    def sets_rows(self) -> int:
        return len(self.sets_refs)

    def dv(self, ws, formula):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        return dv

    # -- MAXIFS over every place a set can be logged ---------------------------------
    def _sources(self, cols, ex_ref, extra=None):
        """MAX of MAXIFS over Start Here baseline, the week sheets (via Sets) and Log."""
        extra = extra or {}
        base_end, sets_end, log_end = BASE_FIRST + BASE_ROWS - 1, self.sets_rows + 1, self.log_rows + 1
        sh, st, lg = cols
        parts = [
            (f"{SH}!${sh}${BASE_FIRST}:${sh}${base_end}", f"{SH}!$A${BASE_FIRST}:$A${base_end}", extra.get("base")),
            (f"Sets!${st}$2:${st}${sets_end}", f"Sets!$B$2:$B${sets_end}", extra.get("sets")),
            (f"Log!${lg}$2:${lg}${log_end}", f"Log!$B$2:$B${log_end}", extra.get("log")),
        ]
        return "MAX(" + ",".join(
            f"_xlfn.MAXIFS({v},{n},{ex_ref}{',' + more if more else ''})" for v, n, more in parts) + ")"

    def _rep_max(self, ex_ref, n):
        return self._sources(("B", "C", "C"), ex_ref, {
            "base": f"{SH}!$C${BASE_FIRST}:$C${BASE_FIRST + BASE_ROWS - 1},\">={n}\"",
            "sets": f'Sets!$D$2:$D${self.sets_rows + 1},">={n}"',
            "log": f'Log!$D$2:$D${self.log_rows + 1},">={n}"'})

    # -- Start Here ----------------------------------------------------------------
    def start_sheet(self):
        ws = self.wb.create_sheet("Start Here", 0)
        ws["A1"] = f"{app_name()} - 1-Year Strength, Mass & Power Program"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = "Yellow cells are yours to fill in."

        ws.cell(3, 1, "Intake").font = SECTION_FONT
        for i, (label, key, default) in enumerate(INTAKE):
            r = 4 + i
            value = self.data.get(key, default)
            if key == "start":
                value = value or next_monday()
            ws.cell(r, 1, label).border = BOX
            c = _input(ws.cell(r, 2), value)
            if key == "start":
                c.number_format = "yyyy-mm-dd"
        keys = [k for _, k, _ in INTAKE]
        for key, formula in (("units", '"lb,kg"'), ("sex", '"Male,Female"'), ("bodyfat_method", self.method_list)):
            self.dv(ws, formula).add(f"B{4 + keys.index(key)}")

        ws.cell(15, 1, "Starting maxes (any movement or variation)").font = SECTION_FONT
        _header(ws, 16, ["Exercise", "Weight", "Reps", "Est. 1RM"])
        dv = self.dv(ws, self.ex_list)
        seed = list(self.data.get("baseline", []))
        defaults = [*MAIN, "Box Squat", "Close-Grip Bench Press", "Romanian Deadlift", "Box Jump"]
        for i in range(BASE_ROWS):
            r = BASE_FIRST + i
            if i < len(seed):
                ex, w, reps = seed[i]
            else:
                ex, w, reps = (defaults[i] if not seed and i < len(defaults) else None), None, None
            _input(ws.cell(r, 1), ex)
            _input(ws.cell(r, 2), w)
            _input(ws.cell(r, 3), reps)
            c = ws.cell(r, 4, _e1rm(f"B{r}", f"C{r}"))
            c.border, c.font = BOX, Font(bold=True)
            dv.add(f"A{r}")
        ws.cell(BASE_FIRST + BASE_ROWS, 1, "Box / broad jump: enter the height or distance as 'Weight' with reps = 1.")

        # PR board (F:K)
        ws["F3"] = "PR board"
        ws["F3"].font = SECTION_FONT
        _header(ws, 4, ["Lift", "Best est. 1RM", "1RM", "3RM", "5RM", "10RM"], first_col=6)
        rm_col = {n: get_column_letter(4 + REP_MAX_COUNTS.index(n)) for n in (1, 3, 5, 10)}
        for i, lift in enumerate(MAIN):
            r, pr = 5 + i, self.row_of[lift]
            ws.cell(r, 6, lift).font = Font(bold=True)
            c = ws.cell(r, 7, f"=PRs!C{pr}")
            c.number_format, c.font = HIDE_ZERO, Font(bold=True, color="1F3A5F")
            for j, n in enumerate((1, 3, 5, 10)):
                ws.cell(r, 8 + j, f"=PRs!{rm_col[n]}{pr}").number_format = HIDE_ZERO
            for col in range(6, 12):
                ws.cell(r, col).border = BOX
        r = 5 + len(MAIN)
        ws.cell(r, 6, "Total (e1RM)").font = Font(bold=True)
        c = ws.cell(r, 7, "=" + "+".join(f"G{5 + i}" for i in range(len(MAIN))))
        c.number_format, c.font = HIDE_ZERO, Font(bold=True)

        # Box jump standards (F:I)
        top = r + 2
        ws.cell(top, 6, "Box jump standards (from your height)").font = SECTION_FONT
        _header(ws, top + 1, ["Level", "Box height", "Target", "Reached?"], first_col=6)
        best_jump = f"PRs!D{self.row_of['Box Jump']}"
        for i, (level, desc, ratio) in enumerate(BOX_JUMP_STANDARDS):
            rr = top + 2 + i
            ws.cell(rr, 6, level).font = Font(bold=True)
            ws.cell(rr, 7, desc)
            ws.cell(rr, 8, f'=IF({UNITS}="kg",19,7.5)' if ratio is None
                    else f'=IF({HEIGHT}="","enter height",ROUND({HEIGHT}*{ratio},1))')
            ws.cell(rr, 9, f'=IF(AND(ISNUMBER(H{rr}),{best_jump}>0),IF({best_jump}>=H{rr},"YES",""),"")').font = PR_FONT
            for col in range(6, 10):
                ws.cell(rr, col).border = BOX
            self.jump_targets[level] = f"{SH}!$H${rr}"
        rr = top + 2 + len(BOX_JUMP_STANDARDS)
        t = self.jump_targets
        ws.cell(rr, 6, "Your best box jump").font = Font(bold=True)
        ws.cell(rr, 8, f"={best_jump}").number_format = HIDE_ZERO
        ws.cell(rr + 1, 6, "Your level").font = Font(bold=True)
        c = ws.cell(rr + 1, 8, f'=IF({best_jump}=0,"Not tested",IF(ISNUMBER({t["Elite"]})*({best_jump}>={t["Elite"]}),'
                               f'"Elite",IF(ISNUMBER({t["Proficient"]})*({best_jump}>={t["Proficient"]}),"Proficient",'
                               f'IF(ISNUMBER({t["Intermediate"]})*({best_jump}>={t["Intermediate"]}),"Intermediate",'
                               f'IF({best_jump}>={t["Beginner"]},"Beginner","Below beginner")))))')
        c.font = Font(bold=True, color="1F3A5F")
        ws.cell(rr + 2, 6, "1 step = beginner, above knee = intermediate, chest = proficient, head height = elite. "
                           "Same unit as your height.")

        # Prilepin's chart (F:I)
        top = rr + 4
        ws.cell(top, 6, "Prilepin's chart").font = SECTION_FONT
        _header(ws, top + 1, ["% of max", "Reps per set", "Optimal total", "Total range"], first_col=6)
        for i, (label, z) in enumerate(zip(("Under 70%", "70-80%", "80-90%", "90%+"), ZONES)):
            for col, val in enumerate((label, f"{z.reps_per_set[0]}-{z.reps_per_set[1]}", z.optimal_total,
                                       f"{z.total_range[0]}-{z.total_range[1]}"), start=6):
                c = ws.cell(top + 2 + i, col, val)
                c.border, c.alignment = BOX, Alignment(horizontal="center")
        ws.cell(top + 6, 6, "Every loaded session's sets x reps come from this chart.")

        # Week links (F:M) and instructions
        top += 8
        ws.cell(top, 6, "Go to week").font = SECTION_FONT
        for week in range(0, WEEKS + 1):
            c = _link(ws.cell(top + 1 + week // 8, 6 + week % 8), week_sheet(week), week_sheet(week))
            phase = next(s.phase for s in self.sessions if s.week == week)
            c.fill = PatternFill("solid", fgColor=PHASE_FILLS[phase])
        top += 3 + WEEKS // 8
        for i, line in enumerate(INSTRUCTIONS):
            ws.cell(top + i, 6, line).font = SECTION_FONT if i == 0 else Font()
        _widths(ws, [30, 14, 8, 11, 3, 16, 14, 12, 10, 9, 9, 9, 9])

    # -- Week sheets -------------------------------------------------------------------
    def week_sheets(self):
        headers = ["Day", "Date", "Exercise", "Sets", "Reps", "% TM", "Training max", "Weight / target",
                   "Actual (weight, height or distance)", "Actual reps", "Est. 1RM", "PR", "Notes", "Guidance"]
        bodyweights = self.data.get("bodyweight", {})
        for week in range(0, WEEKS + 1):
            ws = self.wb.create_sheet(week_sheet(week))
            sessions = [s for s in self.sessions if s.week == week]
            phase = sessions[0].phase
            ws.sheet_properties.tabColor = PHASE_FILLS[phase]
            ws["A1"] = week_label(week)
            ws["A1"].font = Font(bold=True, size=14)
            if week > 0:
                _link(ws["A2"], f"< {week_sheet(week - 1)}", week_sheet(week - 1))
            _link(ws["C2"], "Start Here", "Start Here")
            if week < WEEKS:
                _link(ws["E2"], f"{week_sheet(week + 1)} >", week_sheet(week + 1))
            ws["A3"] = "Bodyweight this week"
            ws["A3"].font = Font(bold=True)
            bw = bodyweights.get(week)
            if bw is None and week == 0:
                bw = f'=IF({START_BW}="","",{START_BW})'
            _input(ws["C3"], bw)
            ws["D3"] = "Weigh in once this week: same morning routine, before eating."
            if week in self.bf_weeks:
                _link(ws["A4"], f"Body fat test due this week (month {self.bf_weeks[week]}) - enter it on the Body sheet",
                      "Body").font = Font(bold=True, color="B00020", underline="single")
            _header(ws, 5, headers)
            ws.freeze_panes = "D6"
            row = WEEK_FIRST_ROW
            for s in sessions:
                offset = (s.week - 1) * 7 + (0, 2, 4)[s.day_index]
                cycle_col = get_column_letter(4 + (s.cycle or 1))  # Maxes: E = cycle 1 ... I = transition
                for p in s.exercises:
                    info = CATALOG.get(p.exercise)
                    pr_row = self.row_of.get(p.exercise)
                    plyo = bool(info and info.category == "plyo")
                    lifting = bool(info and info.category in ("main", "variation"))
                    weight = None
                    if p.is_loaded:
                        weight = f'=IF(OR(G{row}="",F{row}=""),"",MROUND(G{row}*F{row}/100,{ROUNDING}))'
                    elif p.exercise == "Box Jump":
                        t, best = self.jump_targets, f"PRs!D{pr_row}"
                        weight = (f'=IFERROR(IF({best}<{t["Beginner"]},{t["Beginner"]},IF({best}<{t["Intermediate"]},'
                                  f'{t["Intermediate"]},IF({best}<{t["Proficient"]},{t["Proficient"]},'
                                  f'IF({best}<{t["Elite"]},{t["Elite"]},"Elite!")))),"")')
                    if plyo:
                        pr = f'=IF(I{row}="","",IF(I{row}>=PRs!D{pr_row},"PR!",""))'
                    elif pr_row:
                        pr = f'=IF(K{row}="","",IF(K{row}>=PRs!C{pr_row},"PR!",""))'
                    else:
                        pr = None
                    values = [s.day, f"={START}+{offset}", p.exercise, p.sets or None, p.reps, p.percent,
                              f"=Maxes!{cycle_col}{pr_row}" if lifting else None, weight, None, None,
                              None if plyo else _e1rm(f"I{row}", f"J{row}"), pr, None, p.note]
                    for col, val in enumerate(values, start=1):
                        c = ws.cell(row, col, val)
                        c.border = BOX
                        c.alignment = Alignment(vertical="center", wrap_text=col == 14)
                    ws.cell(row, 1).fill = DAY_FILL
                    ws.cell(row, 2).number_format = "ddd mmm d"
                    ws.cell(row, 7).number_format = "0"
                    ws.cell(row, 8).font = Font(bold=True)
                    ws.cell(row, 12).font = PR_FONT
                    for col in (9, 10, 13):
                        ws.cell(row, col).fill = INPUT_FILL
                    if p.kind in ("main", "variation", "test"):
                        ws.cell(row, 3).font = Font(bold=True)
                    elif plyo:
                        ws.cell(row, 3).font = Font(bold=True, color="2E7D32")
                    self.sets_refs.append((ws.title, row))
                    row += 1
            _widths(ws, [17, 11, 24, 5, 6, 6, 9, 10, 13, 8, 9, 5, 22, 62])
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True

    def sets_sheet(self):
        """Hidden mirror of every week-sheet row so PR formulas can scan one range."""
        ws = self.wb.create_sheet("Sets")
        ws.append(["Date", "Exercise", "Weight", "Reps", "Est. 1RM"])
        for i, (sheet, r) in enumerate(self.sets_refs, start=2):
            q = f"'{sheet}'!"
            ws.cell(i, 1, f"={q}B{r}")
            ws.cell(i, 2, f"={q}C{r}")
            ws.cell(i, 3, f'=IF({q}I{r}="","",{q}I{r})')
            ws.cell(i, 4, f'=IF({q}I{r}="","",IF({q}J{r}="",1,{q}J{r}))')  # jumps count as 1 rep
            ws.cell(i, 5, _e1rm(f"C{i}", f"D{i}"))
        ws.sheet_state = "hidden"

    # -- Log -------------------------------------------------------------------------
    def log_sheet(self):
        ws = self.wb.create_sheet("Log")
        _header(ws, 1, ["Date", "Exercise", "Weight / height", "Reps", "Est. 1RM", "Notes"])
        ws.freeze_panes = "A2"
        entries = list(self.data.get("log", []))
        for i in range(self.log_rows):
            r = 2 + i
            d, ex, w, reps, note = entries[i] if i < len(entries) else (None,) * 5
            for col, val in enumerate((d, ex, w, reps), start=1):
                _input(ws.cell(r, col), val)
            ws.cell(r, 1).number_format = "yyyy-mm-dd"
            ws.cell(r, 5, _e1rm(f"C{r}", f"D{r}")).border = BOX
            _input(ws.cell(r, 6), note or None)
        self.dv(ws, self.ex_list).add(f"B2:B{self.log_rows + 1}")
        _widths(ws, [12, 26, 14, 7, 10, 40])

    # -- PRs -------------------------------------------------------------------------
    def prs_sheet(self):
        ws = self.wb.create_sheet("PRs")
        _header(ws, 1, ["Exercise", "Type", "Best est. 1RM", *[f"{n}RM" for n in REP_MAX_COUNTS]])
        ws.freeze_panes = "C2"
        for name, r in self.row_of.items():
            info = CATALOG[name]
            ws.cell(r, 1, name).font = Font(bold=info.category == "main")
            ws.cell(r, 2, info.category)
            if info.category != "plyo":
                c = ws.cell(r, 3, "=" + self._sources(("D", "E", "E"), f"$A{r}"))
                c.number_format, c.font = HIDE_ZERO, Font(bold=True)
            for j, n in enumerate(REP_MAX_COUNTS):
                if info.category == "plyo" and n > 1:
                    break
                ws.cell(r, 4 + j, "=" + self._rep_max(f"$A{r}", n)).number_format = HIDE_ZERO
        ws.cell(len(self.row_of) + 3, 1,
                "NRM = heaviest weight lifted for at least N reps (a 5-rep set also counts toward your 3RM). "
                "Jumps: the 1RM column is your best height / distance. Pulls from Start Here, every week and Log.")
        _widths(ws, [26, 10, 13] + [8] * len(REP_MAX_COUNTS))

    # -- Maxes -----------------------------------------------------------------------
    def maxes_sheet(self):
        ws = self.wb.create_sheet("Maxes")
        labels = [f"Cycle {c}\n(weeks {(c - 1) * 12 + 1}-{c * 12})" for c in range(1, CYCLES + 1)]
        _header(ws, 1, ["Exercise", "Type", "Parent lift", "Ratio", *labels, "Transition\n(weeks 49-52)"])
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "E2"
        for name, r in self.row_of.items():
            info = CATALOG[name]
            ws.cell(r, 1, name).font = Font(bold=info.category == "main")
            ws.cell(r, 2, info.category)
            ws.cell(r, 3, info.parent)
            ws.cell(r, 4, info.ratio)
            if info.category not in ("main", "variation"):
                continue
            for k in range(1, CYCLES + 2):
                col = 4 + k
                cs = f"{START}+{(k - 1) * CYCLE_WEEKS * 7}"
                own = self._sources(("D", "E", "E"), f"$A{r}", {
                    "sets": f'Sets!$A$2:$A${self.sets_rows + 1},"<"&({cs})',
                    "log": f'Log!$A$2:$A${self.log_rows + 1},"<"&({cs})'})
                if info.parent:
                    parent = f"{get_column_letter(col)}{self.row_of[info.parent]}"
                    formula = f'=IF({own}>0,ROUND({own},1),IF(N({parent})>0,ROUND({parent}*$D{r},1),""))'
                else:
                    formula = f'=IF({own}>0,ROUND({own},1),"")'
                ws.cell(r, col, formula).number_format = "0"
        ws.conditional_formatting.add(
            f"E2:I{1 + len(self.row_of)}",
            FormulaRule(formula=['AND($B2="variation",E2<>"",PRs!$C2=0)'], font=Font(italic=True, color="888888")))
        ws.cell(len(self.row_of) + 3, 1,
                "Training max = best estimated 1RM logged before the cycle starts. Grey italics = estimated "
                "from the parent lift x ratio (no sets logged yet).")
        _widths(ws, [26, 10, 15, 7, 12, 12, 12, 12, 13])

    # -- Body ------------------------------------------------------------------------
    def body_sheet(self):
        ws = self.wb.create_sheet("Body")
        ws["A1"] = "Bodyweight - weekly"
        ws["A1"].font = SECTION_FONT
        ws["A2"] = "Entered at the top of each week sheet."
        _header(ws, 3, ["Week", "Week of", "Bodyweight"])
        for week in range(0, WEEKS + 1):
            r = 4 + week
            _link(ws.cell(r, 1), week, week_sheet(week)).border = BOX
            c = ws.cell(r, 2, f"={START}+{(week - 1) * 7}")
            c.number_format, c.border = "yyyy-mm-dd", BOX
            q = f"'{week_sheet(week)}'!C3"
            ws.cell(r, 3, f'=IF({q}="","",{q})').border = BOX

        ws["E1"] = "Body fat - monthly"
        ws["E1"].font = SECTION_FONT
        ws["E2"] = "Same method, same place, same time of day every month. How-to guide below."
        _header(ws, 3, ["Month", "Test around", "Date tested", "Body fat %", "Method", "Weight at test",
                        "Lean mass", "Fat mass"], first_col=5)
        dv = self.dv(ws, self.method_list)
        fats = self.data.get("bodyfat", {})
        for month in range(0, 13):
            r = 4 + month
            ws.cell(r, 5, "Start" if month == 0 else month).border = BOX
            if month == 0:  # from the intake
                ws.cell(r, 6, f"={START}").number_format = "yyyy-mm-dd"
                ws.cell(r, 8, f'=IF({START_BF}="","",{START_BF})')
                ws.cell(r, 9, f'=IF({BF_METHOD}="","",{BF_METHOD})')
                ws.cell(r, 10, f'=IF({START_BW}="","",{START_BW})')
                for col in range(6, 11):
                    ws.cell(r, col).border = BOX
            else:
                d, pct, method, weight = fats.get(month, (None,) * 4)
                c = ws.cell(r, 6, f"={START}+{(bodyfat_week(month) - 1) * 7}")
                c.number_format, c.border = "yyyy-mm-dd", BOX
                _input(ws.cell(r, 7), d).number_format = "yyyy-mm-dd"
                _input(ws.cell(r, 8), pct)
                _input(ws.cell(r, 9), method)
                _input(ws.cell(r, 10), weight)
                dv.add(f"I{r}")
            ws.cell(r, 11, f'=IF(OR(H{r}="",J{r}=""),"",ROUND(J{r}*(1-H{r}/100),1))').border = BOX
            ws.cell(r, 12, f'=IF(OR(H{r}="",J{r}=""),"",ROUND(J{r}*H{r}/100,1))').border = BOX
        headings = {"How to measure body fat once a month", "For every method"}
        for i, line in enumerate(BODYFAT_GUIDE.splitlines()):
            c = ws.cell(19 + i, 5, line)
            if line in headings or line.endswith(")"):
                c.font = Font(bold=True)
        _widths(ws, [7, 12, 12, 3, 7, 12, 12, 11, 22, 13, 11, 10])

    # -- Progress --------------------------------------------------------------------
    def progress_sheet(self):
        ws = self.wb.create_sheet("Progress")
        ws["A1"] = "Progress - strength vs bodyweight and body fat"
        ws["A1"].font = Font(bold=True, size=14)
        n = len(MAIN)
        _header(ws, 3, ["Checkpoint", "Week", *MAIN, "Bodyweight", "Body fat %"])
        points = [("Start", 1)] + [(f"After cycle {c}", c * CYCLE_WEEKS + 1) for c in range(1, CYCLES + 1)]
        for i, (label, week) in enumerate(points):
            r = 4 + i
            ws.cell(r, 1, label).border = BOX
            ws.cell(r, 2, week).border = BOX
            for j, lift in enumerate(MAIN):
                ref = f"Maxes!{get_column_letter(5 + i)}{self.row_of[lift]}"
                c = ws.cell(r, 3 + j, f'=IF({ref}="",NA(),{ref})')
                c.number_format, c.border = "0", BOX
            # latest weekly bodyweight at or before this checkpoint, else the intake value
            bw = (f'=IFERROR(LOOKUP(2,1/((Body!$A$4:$A$56<=B{r})*(Body!$C$4:$C$56<>"")),Body!$C$4:$C$56),'
                  f'IF({START_BW}="",NA(),{START_BW}))')
            ws.cell(r, 3 + n, bw).border = BOX
            month = max(0, min(12, (week - 1) * 12 // 52))
            bf = (f'=IFERROR(LOOKUP(2,1/((ROW(Body!$H$4:$H$16)-4<={month})*(Body!$H$4:$H$16<>"")),'
                  f'Body!$H$4:$H$16),NA())')
            ws.cell(r, 4 + n, bf).border = BOX
        ws.cell(10, 1, "Strength = training max at each checkpoint (best estimated 1RM so far). "
                       "#N/A just means no data yet.")

        strength = LineChart()
        strength.title = "Strength (est. 1RM) vs bodyweight"
        strength.y_axis.title = "Estimated 1RM"
        strength.add_data(Reference(ws, min_col=3, max_col=2 + n, min_row=3, max_row=8), titles_from_data=True)
        strength.set_categories(Reference(ws, min_col=1, min_row=4, max_row=8))
        body = LineChart()
        body.add_data(Reference(ws, min_col=3 + n, min_row=3, max_row=8), titles_from_data=True)
        body.y_axis.axId = 200
        body.y_axis.title = "Bodyweight"
        body.y_axis.crosses = "max"
        strength += body
        strength.height, strength.width = 9, 18
        ws.add_chart(strength, "A12")

        body_ws = self.wb["Body"]
        bw = LineChart()
        bw.title = "Bodyweight (weekly)"
        bw.y_axis.title = "Bodyweight"
        bw.x_axis.title = "Week"
        bw.add_data(Reference(body_ws, min_col=3, min_row=3, max_row=4 + WEEKS), titles_from_data=True)
        bw.set_categories(Reference(body_ws, min_col=1, min_row=4, max_row=4 + WEEKS))
        bw.height, bw.width = 9, 18
        ws.add_chart(bw, "L3")

        bf = LineChart()
        bf.title = "Body fat % and lean mass (monthly)"
        bf.y_axis.title = "Body fat %"
        bf.x_axis.title = "Month"
        bf.add_data(Reference(body_ws, min_col=8, min_row=3, max_row=16), titles_from_data=True)
        bf.set_categories(Reference(body_ws, min_col=5, min_row=4, max_row=16))
        lean = LineChart()
        lean.add_data(Reference(body_ws, min_col=11, min_row=3, max_row=16), titles_from_data=True)
        lean.y_axis.axId = 300
        lean.y_axis.title = "Lean mass"
        lean.y_axis.crosses = "max"
        bf += lean
        bf.height, bf.width = 9, 18
        ws.add_chart(bf, "L22")
        _widths(ws, [16, 7] + [11] * n + [12, 11])


def build_workbook(path: str | Path, data: dict | None = None) -> Path:
    """Write the workbook to `path`.

    `data` optionally pre-fills it:
      intake keys: name, units, increment, start, height, bodyweight_start,
                   bodyfat_start, bodyfat_method, age, sex
      baseline:    [(exercise, weight, reps), ...]           -> Start Here
      log:         [(date, exercise, weight, reps, note)]    -> Log
      bodyweight:  {week: weight}                            -> week sheets
      bodyfat:     {month: (date, percent, method, weight)}  -> Body
    """
    book = _Book(data)
    book.start_sheet()
    book.week_sheets()   # must run before formulas that scan the Sets mirror
    book.log_sheet()
    book.prs_sheet()
    book.maxes_sheet()
    book.body_sheet()
    book.progress_sheet()
    book.sets_sheet()
    wb = book.wb
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    wb.active = 0
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
